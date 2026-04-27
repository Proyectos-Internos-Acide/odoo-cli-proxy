#!/usr/bin/env python3
"""Fix price_extra=0 on all variant products created by import_cotizacion_2026.py.
Self-contained: re-parses Excel and updates PTAVs directly."""

import xmlrpc.client
import openpyxl
import re
from collections import defaultdict

URL = 'https://machupicchu-afdestiny-production.odoo.com/'
DB = 'machupicchu-afdestiny-production'
USER = 'ayfdestinyeirl@gmail.com'
PW = 'L^yUbD2^+p^2h.#'
EXCEL = '/home/gonzalo/Descargas/COTIZACION 2026.xlsx'

common = xmlrpc.client.ServerProxy(f'{URL}xmlrpc/2/common')
uid = common.authenticate(DB, USER, PW, {})
models = xmlrpc.client.ServerProxy(f'{URL}xmlrpc/2/object')

def ex(model, method, *a, **kw):
    return models.execute_kw(DB, uid, PW, model, method, *a, **kw)

wb = openpyxl.load_workbook(EXCEL, data_only=True)

# ---- Helpers ----
CITY_MAP = {'Mchupicchu': 'Machupicchu'}
def ncity(r): return CITY_MAP.get((r or '').strip(), (r or '').strip())

def room_type(th):
    if not th: return None
    m = re.search(r'\s*-\s*(Simple|Doble o Matrimonial|Tripe|Triple|Triples|Familiar\s*)\s*$', th, re.I)
    if not m: return None
    rt = m.group(1).strip()
    if rt in ('Tripe','Triple','Triples'): return 'Triple'
    if rt.startswith('Familiar'): return 'Familiar'
    return rt

def hotel_from_tipo(th):
    if not th: return None
    m = re.search(r'^(.+?)\s*-\s*(?:Simple|Doble o Matrimonial|Tripe|Triple|Familiar)', th, re.I)
    return m.group(1).strip() if m else None

def pax_label(name):
    m = re.search(r'x\s*(\d+(?:\s*-\s*\d+)?)\s*(?:pax)?$', name.strip(), re.I)
    if m: return f"{m.group(1).replace(' ','')} pax"
    m = re.search(r'\((\+?\d+(?:\s*-?\s*\d+)?)\s*pax\)', name, re.I)
    if m: return f"{m.group(1).replace(' ','')} pax"
    m = re.search(r'-\s*(\d+(?:\s*-\s*\d+)?)\s*pax$', name.strip(), re.I)
    if m: return f"{m.group(1).replace(' ','')} pax"
    return None

def base_name(name):
    c = re.sub(r'\s*x\s*\d+(?:\s*-\s*\d+)?\s*(?:pax)?\s*$', '', name.strip(), flags=re.I)
    c = re.sub(r'\s*\(\+?\d+(?:\s*-?\s*\d+)?\s*pax\)\s*$', '', c, flags=re.I)
    c = re.sub(r'\s*-\s*\d+(?:\s*-\s*\d+)?\s*pax\s*$', '', c, flags=re.I)
    return c.strip()

TYPOS = {'Valle sagrado - Pisac - Ollantao': 'Valle sagrado - Pisac - Ollanta'}
RT_ORDER = ['Simple', 'Doble o Matrimonial', 'Triple', 'Familiar']
PAX_STD = ['1 pax','2 pax','3 pax','4 pax','5-6 pax','7-8 pax','9-10 pax']

# ---- Build price_map: product_name -> {variant_name: price} ----
price_map = {}

# Hotels
for stars in ['2 estrellas','3 estrellas','4 estrellas','5 estrellas']:
    ws = wb[f'Hotel {stars}']
    sn = stars.split()[0]
    htls = defaultdict(lambda: {'rooms': []})
    for row in ws.iter_rows(min_row=2, values_only=True):
        th, hn_col = row[3], (row[4] or '').strip() if row[4] else None
        pn, pp = row[5], row[6]
        if not th: continue
        rt = room_type(th)
        if not rt: continue
        hn = hotel_from_tipo(th) or hn_col
        if not hn or hn.lower() in ('arbnb','none',''): continue
        city = ncity(row[1]) if row[1] else None
        ppv = pp if pp and pp > 0 else 0
        pnv = pn if pn and pn > 0 else 0
        if ppv == 0 and pnv == 0: continue
        htls[(hn, sn, city or 'Otro')]['rooms'].append((rt, ppv))
    # Detect cross-city collisions
    name_cities = defaultdict(set)
    for (hn, sn2, city) in htls: name_cities[hn].add(city)
    for (hn, sn2, city), data in htls.items():
        dname = f"{hn} ({city})" if len(name_cities[hn]) > 1 else hn
        seen = {}
        for rt2, pp2 in data['rooms']:
            if rt2 not in seen: seen[rt2] = pp2
        variants = {rt2: round(seen[rt2], 2) for rt2 in RT_ORDER if rt2 in seen and seen[rt2] > 0}
        if variants: price_map[dname] = variants

# Traslados
ws = wb['Traslados']
routes = defaultdict(list)
for row in ws.iter_rows(min_row=2, values_only=True):
    route = (row[2] or '').strip()
    pr = (row[4] or '').strip()
    price = row[5]
    if not route or not pr or not price or price == 0: continue
    pl = pax_label(pr)
    if not pl:
        m = re.search(r'(\d+(?:\s*-\s*\d+)?)\s*pax', pr, re.I)
        if m: pl = f"{m.group(1).replace(' ','')} pax"
        else: continue
    routes[route].append((pl, round(price, 2)))
for rn, vlist in routes.items():
    seen = {}
    for p, pr in vlist:
        if p not in seen: seen[p] = pr
    ordered = {p: seen[p] for p in PAX_STD if p in seen}
    for p in seen:
        if p not in PAX_STD: ordered[p] = seen[p]
    if ordered: price_map[rn] = ordered

# Tours compartidos (grouped only)
ws = wb['Tours Compartidos']
skip = {'OOOOOO','','o','p','q'}
tc_grouped = defaultdict(list)
for row in ws.iter_rows(min_row=2, values_only=True):
    name = (row[2] or '').strip()
    price = row[6]
    if not name or name in skip or not price or price == 0: continue
    pl = pax_label(name)
    if pl:
        bn = base_name(name)
        tc_grouped[bn].append((pl, round(price, 2)))
for bn, vlist in tc_grouped.items():
    price_map[bn] = {p: pr for p, pr in vlist if pr > 0}

# Tours privados
ws = wb['Tours Privados']
tp_grouped = defaultdict(list)
for row in ws.iter_rows(min_row=2, values_only=True):
    name = (row[2] or '').strip()
    price = row[6]
    if not name or not price or price == 0 or isinstance(price, str): continue
    pl = pax_label(name)
    bn = base_name(name)
    bn = TYPOS.get(bn, bn)
    if pl and bn: tp_grouped[bn].append((pl, round(price, 2)))
for bn, vlist in tp_grouped.items():
    price_map[bn] = {p: pr for p, pr in vlist if pr > 0}

print(f"Parsed {len(price_map)} products with variant prices from Excel")

# ---- Find Agencia variant products and update price_extra ----
cats = ex('product.category', 'search', [[['name', 'ilike', 'Agencia']]])
tmpls = ex('product.template', 'search_read',
    [['&', ['categ_id', 'in', cats], ['attribute_line_ids', '!=', False]]],
    {'fields': ['id', 'name', 'attribute_line_ids']})

updated = 0
skipped = 0
not_found = []

for t in tmpls:
    name = t['name']
    if name not in price_map:
        not_found.append(name)
        continue
    expected = price_map[name]
    for al_id in t['attribute_line_ids']:
        ptavs = ex('product.template.attribute.value', 'search_read',
            [[['attribute_line_id', '=', al_id]]],
            {'fields': ['id', 'name', 'price_extra']})
        for ptav in ptavs:
            vname = ptav['name']
            if vname in expected:
                new_price = round(expected[vname], 2)
                if abs(ptav['price_extra'] - new_price) > 0.001:
                    ex('product.template.attribute.value', 'write',
                       [[ptav['id']], {'price_extra': new_price}])
                    print(f"  [{t['id']}] {name} / {vname}: {ptav['price_extra']} -> {new_price}")
                    updated += 1
                else:
                    skipped += 1

print(f"\nDone! Updated: {updated}, Already correct: {skipped}, Not in Excel: {len(not_found)}")
if not_found:
    print("Not found in Excel:")
    for n in not_found:
        print(f"  - {n}")
